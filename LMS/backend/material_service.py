import csv
import hashlib
import io
import logging
import os
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

from bson import ObjectId
from PIL import Image

from mongodb import mongo_db
from vector_store import index_resource, clear_search_cache

logger = logging.getLogger("lms.materials")
ALLOWED = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
    ".doc": {"application/msword", "application/octet-stream"},
    ".pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation", "application/octet-stream"},
    ".ppt": {"application/vnd.ms-powerpoint", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".webp": {"image/webp"},
    ".csv": {"text/csv", "application/csv", "application/vnd.ms-excel", "text/plain", "application/octet-stream"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/octet-stream"},
}
MAX_EXTRACTED_CHARS = int(os.getenv("MAX_EXTRACTED_CHARS", "120000"))
MAX_ARCHIVE_UNCOMPRESSED_BYTES = max(1, int(os.getenv("MAX_ARCHIVE_UNCOMPRESSED_MB", "100"))) * 1024 * 1024
MAX_ARCHIVE_ENTRIES = max(1, int(os.getenv("MAX_ARCHIVE_ENTRIES", "5000")))
MAX_IMAGE_PIXELS = max(1, int(os.getenv("MAX_IMAGE_PIXELS", "40000000")))
MAX_ARCHIVE_MEMBER_BYTES = max(1, int(os.getenv("MAX_ARCHIVE_MEMBER_MB", "25"))) * 1024 * 1024
MAX_COMPRESSION_RATIO = max(1, int(os.getenv("MAX_ARCHIVE_COMPRESSION_RATIO", "200")))


def safe_filename(name: str) -> str:
    # Normalize both Windows and POSIX paths before using the final path component.
    value = str(name or "").replace("\\", "/").replace("\x00", "").split("/")[-1].strip()
    if not value or value in {".", ".."} or len(value) > 180:
        raise ValueError("Invalid filename")
    return value


def _validate_office_archive(data: bytes, suffix: str) -> None:
    required = {
        ".docx": "word/document.xml",
        ".pptx": "ppt/presentation.xml",
        ".xlsx": "xl/workbook.xml",
    }
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("Malformed Office file")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Office archive contains too many entries")
            if required[suffix] not in {member.filename for member in members}:
                raise ValueError("Malformed Office file")
            total_uncompressed = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise ValueError("Encrypted Office files are not supported")
                if member.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise ValueError("Office archive contains an oversized entry")
                total_uncompressed += member.file_size
                if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise ValueError("Office archive expands beyond the safe processing limit")
                if member.file_size > 0 and member.compress_size == 0:
                    raise ValueError("Office archive has an invalid compression ratio")
                if member.compress_size > 0 and member.file_size / member.compress_size > MAX_COMPRESSION_RATIO:
                    raise ValueError("Office archive compression ratio exceeds the safe limit")
    except zipfile.BadZipFile as exc:
        raise ValueError("Malformed Office file") from exc


def validate_file(name: str, content_type: str, data: bytes) -> str:
    filename = safe_filename(name)
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("Unsupported file format")
    if content_type and content_type not in ALLOWED[suffix]:
        raise ValueError("MIME type does not match the file format")
    if not data:
        raise ValueError("Empty file")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF signature")
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                    raise ValueError("Image dimensions exceed the safe processing limit")
                image.verify()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Malformed image file") from exc
    elif suffix in {".docx", ".pptx", ".xlsx"}:
        _validate_office_archive(data, suffix)
    elif suffix in {".doc", ".ppt"}:
        if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError("Invalid legacy Office file signature")
    elif suffix in {".txt", ".md", ".csv"}:
        if b"\x00" in data[:4096]:
            raise ValueError("Binary content is not valid for this text format")
    return suffix


def extract_content(data: bytes, suffix: str) -> tuple[str, dict]:
    meta = {}
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append("[Page " + str(number) + "]\n" + text)
        meta["page_count"] = len(reader.pages)
        return "\n\n".join(parts)[:MAX_EXTRACTED_CHARS], meta
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)[:MAX_EXTRACTED_CHARS], meta
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        parts = []
        for number, slide in enumerate(presentation.slides, 1):
            text = " ".join(shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
            if text:
                parts.append("[Slide " + str(number) + "] " + text)
        meta["slide_count"] = len(presentation.slides)
        return "\n".join(parts)[:MAX_EXTRACTED_CHARS], meta
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")[:MAX_EXTRACTED_CHARS], meta
    if suffix == ".csv":
        rows = []
        reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
        for row in reader:
            rows.append(" | ".join(row))
            if sum(len(x) for x in rows) >= MAX_EXTRACTED_CHARS:
                break
        return "\n".join(rows)[:MAX_EXTRACTED_CHARS], meta
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        try:
            for sheet in workbook.worksheets:
                rows.append("[Sheet " + sheet.title + "]")
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append(" | ".join(values))
                    if sum(len(x) for x in rows) >= MAX_EXTRACTED_CHARS:
                        break
                if sum(len(x) for x in rows) >= MAX_EXTRACTED_CHARS:
                    break
        finally:
            workbook.close()
        return "\n".join(rows)[:MAX_EXTRACTED_CHARS], meta
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            import pytesseract
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                    raise ValueError("Image dimensions exceed the safe processing limit")
                text = pytesseract.image_to_string(image).strip()
            meta["ocr_status"] = "READY" if text else "EMPTY"
            return text[:MAX_EXTRACTED_CHARS], meta
        except Exception as exc:
            logger.warning("OCR unavailable or failed: %s", type(exc).__name__)
            meta["ocr_status"] = "UNAVAILABLE"
            return "", meta
    if suffix in {".doc", ".ppt"}:
        command = "antiword" if suffix == ".doc" else "catppt"
        executable = shutil.which(command)
        if not executable:
            raise ValueError("Legacy Office format is stored safely but cannot be indexed on this deployment; convert it to DOCX or PPTX")
        process = subprocess.run([executable, "-"], input=data, capture_output=True, timeout=10, check=False)
        if process.returncode != 0:
            raise ValueError("Legacy Office extraction failed")
        return process.stdout.decode("utf-8", errors="replace")[:MAX_EXTRACTED_CHARS], meta
    raise ValueError("Unsupported file format")


def prepare_resource(file_id: str):
    resource = mongo_db.resources.find_one({"_id": ObjectId(file_id)})
    if not resource:
        return False
    mongo_db.resources.update_one(
        {"_id": resource["_id"]},
        {"$set": {"processing_status": "PROCESSING", "extraction_status": "PROCESSING", "indexing_status": "PENDING", "updated_at": datetime.utcnow()}},
    )
    try:
        content, metadata = extract_content(resource["file"], resource["extension"])
        if not content.strip() and resource["extension"] not in {".png", ".jpg", ".jpeg", ".webp"}:
            raise ValueError("No extractable text found")
        mongo_db.resources.update_one(
            {"_id": resource["_id"]},
            {"$set": {"content": content, "extraction_status": "READY", "processing_status": "PROCESSING", "indexing_status": "PROCESSING", "updated_at": datetime.utcnow(), **metadata}},
        )
    except Exception as exc:
        logger.exception("Material extraction failed for %s", file_id)
        mongo_db.resources.update_one(
            {"_id": resource["_id"]},
            {"$set": {"processing_status": "FAILED", "extraction_status": "FAILED", "indexing_status": "FAILED", "error_message": str(exc)[:1000], "updated_at": datetime.utcnow()}},
        )
        return False

    # Extraction and semantic indexing are separate reliability boundaries. If
    # the optional vector model/store is unavailable, retain the extracted text
    # and expose the resource to the lexical-search fallback instead of marking
    # a valid upload as failed.
    indexing_error = None
    try:
        indexed = index_resource({
            "id": "resource-" + str(resource["_id"]),
            "title": resource.get("title") or resource.get("filename", "Learning material"),
            "content": content,
            "type": "resource",
            "course_id": resource["course_id"],
        })
    except Exception as exc:
        indexed = False
        indexing_error = type(exc).__name__
        logger.warning("Semantic indexing unavailable for resource %s; lexical fallback remains available (%s)", file_id, indexing_error)

    final_fields = {
        "processing_status": "READY",
        "extraction_status": "READY",
        "indexing_status": "READY" if indexed else "LEXICAL_ONLY",
        "updated_at": datetime.utcnow(),
    }
    if indexing_error:
        final_fields["indexing_error"] = "Semantic indexing unavailable; keyword search remains available."
    else:
        final_fields["indexing_error"] = None
    mongo_db.resources.update_one({"_id": resource["_id"]}, {"$set": final_fields})
    clear_search_cache()
    return True


def duplicate_hash(course_id: int, sha256: str):
    return mongo_db.resources.find_one(
        {"course_id": course_id, "sha256": sha256, "processing_status": {"$ne": "FAILED"}},
        {"_id": 1, "filename": 1},
    )


def build_metadata(file, data: bytes, course_id: int, teacher_id: int, title: str):
    filename = safe_filename(file.filename or "learning-material")
    suffix = validate_file(filename, file.content_type or "", data)
    return {
        "course_id": course_id,
        "teacher_id": teacher_id,
        "title": (title or Path(filename).stem).strip()[:200] or Path(filename).stem,
        "filename": filename,
        "content_type": file.content_type or "application/octet-stream",
        "extension": suffix,
        "size": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "content": "",
        "processing_status": "UPLOADED",
        "extraction_status": "PENDING",
        "indexing_status": "PENDING",
        "created_at": datetime.utcnow(),
        "updated_at": datetime.utcnow(),
        "file": data,
    }
