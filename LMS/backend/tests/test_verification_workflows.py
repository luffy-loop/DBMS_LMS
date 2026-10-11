"""Isolated API-level regression checks for LMS material and assessment workflows.

All rows are synthetic and stored in temporary SQLite / in-memory Mongo-like fixtures.
Nothing in this module writes to Render or any real student/teacher account.
"""
import os
os.environ["ENVIRONMENT"] = "test"
os.environ["RUN_DB_SETUP"] = "false"
os.environ["ENABLE_RESOURCE_VECTOR_SEARCH"] = "false"

import time
from datetime import datetime
from types import SimpleNamespace

from bson import ObjectId
from fastapi.testclient import TestClient
from jose import jwt
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import main
import material_service
from auth import alg, key
from database import get_db
from models import (
    Assignment, AssessmentQuestion, AssessmentQuestionOption, AssessmentQuestionRubric,
    AuditLog, Course, Enrollment, StudentQuestionAnswer, Submission, User,
)


class FakeCursor:
    def __init__(self, rows):
        self.rows = list(rows)

    def sort(self, field, direction=1):
        self.rows.sort(key=lambda row: row.get(field) or datetime.min, reverse=direction < 0)
        return self

    def limit(self, count):
        self.rows = self.rows[:count]
        return self

    def __iter__(self):
        return iter(self.rows)


class FakeResources:
    """Small Mongo collection test double supporting predicates used by these routes."""
    def __init__(self):
        self.rows = []

    @staticmethod
    def matches(row, query):
        for key_name, expected in (query or {}).items():
            if key_name == "$or":
                if not any(FakeResources.matches(row, item) for item in expected):
                    return False
                continue
            actual = row.get(key_name, None)
            if isinstance(expected, dict):
                if "$exists" in expected and ((key_name in row) != bool(expected["$exists"])):
                    return False
                if "$ne" in expected and actual == expected["$ne"]:
                    return False
                if "$in" in expected and actual not in expected["$in"]:
                    return False
                if "$lt" in expected and not (actual is not None and actual < expected["$lt"]):
                    return False
            elif actual != expected:
                return False
        return True

    def find_one(self, query, projection=None):
        return next((row for row in self.rows if self.matches(row, query)), None)

    def find(self, query=None, projection=None):
        return FakeCursor([row for row in self.rows if self.matches(row, query or {})])

    def insert_one(self, document):
        stored = dict(document)
        stored["_id"] = ObjectId()
        self.rows.append(stored)
        return SimpleNamespace(inserted_id=stored["_id"])

    def update_one(self, query, update):
        row = self.find_one(query)
        if row is None:
            return SimpleNamespace(modified_count=0)
        row.update(update.get("$set", {}))
        for field in update.get("$unset", {}):
            row.pop(field, None)
        return SimpleNamespace(modified_count=1)

    def find_one_and_update(self, query, update, return_document=None):
        row = self.find_one(query)
        if row is None:
            return None
        row.update(update.get("$set", {}))
        for field, increment in update.get("$inc", {}).items():
            row[field] = row.get(field, 0) + increment
        return row

    def delete_many(self, query):
        old_len = len(self.rows)
        self.rows[:] = [row for row in self.rows if not self.matches(row, query)]
        return SimpleNamespace(deleted_count=old_len - len(self.rows))


class FakeSubmissionFiles:
    def find(self, query=None, projection=None):
        return FakeCursor([])

    def find_one(self, query=None, projection=None):
        return None


async def _skip_background_processing(*args, **kwargs):
    return True


def install_database(monkeypatch):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    tables = [
        User.__table__, Course.__table__, Enrollment.__table__, Assignment.__table__,
        AssessmentQuestion.__table__, AssessmentQuestionOption.__table__, AssessmentQuestionRubric.__table__,
        Submission.__table__, StudentQuestionAnswer.__table__, AuditLog.__table__,
    ]
    for table in tables:
        table.create(engine, checkfirst=True)
    db = Session()
    db.add_all([
        User(id=11, name="Teacher One", email="TEACHER001", password="fixture", role="teacher", section="T1"),
        User(id=12, name="Enrolled Student", email="STUDENT001", password="fixture", role="student", section="S1"),
        User(id=13, name="Unenrolled Student", email="STUDENT002", password="fixture", role="student", section="S2"),
        User(id=14, name="Teacher Two", email="TEACHER002", password="fixture", role="teacher", section="T2"),
        User(id=15, name="Admin", email="ADMIN001", password="fixture", role="admin", section="A1"),
    ])
    db.commit()
    db.add_all([
        Course(id=101, title="Database Systems", description="Relational databases and indexes", teacher_id=11),
        Course(id=102, title="Operating Systems", description="Processes and synchronization", teacher_id=14),
    ])
    db.commit()
    db.add(Enrollment(student_id=12, course_id=101))
    db.commit()

    def override_db():
        yield db

    main.app.dependency_overrides[get_db] = override_db
    fake_resources = FakeResources()
    monkeypatch.setattr(main.mongo_db, "resources", fake_resources)
    monkeypatch.setattr(main.mongo_db, "submission_files", FakeSubmissionFiles())
    monkeypatch.setattr(material_service, "index_resource", lambda item: False)
    monkeypatch.setattr(material_service, "clear_search_cache", lambda: None)
    monkeypatch.setattr(main, "_process_and_notify", _skip_background_processing)
    monkeypatch.setattr(main, "notify_users", lambda *args, **kwargs: 0)
    monkeypatch.setattr(main, "notify_submission_received", lambda *args, **kwargs: None)
    return engine, db, tables, fake_resources


def token(user_id, role, exp=None):
    expires = int(time.time()) + 300 if exp is None else exp
    return jwt.encode({"id": user_id, "role": role, "exp": expires}, key, algorithm=alg)


def headers(user_id, role):
    return {"Authorization": "Bearer " + token(user_id, role)}


def close_database(engine, db, tables):
    main.app.dependency_overrides.pop(get_db, None)
    db.close()
    for table in reversed(tables):
        table.drop(engine, checkfirst=True)
    engine.dispose()


def test_security_http_matrix_reports_expected_and_actual_statuses(monkeypatch):
    engine, db, tables, _ = install_database(monkeypatch)
    try:
        client = TestClient(main.app)
        expired = token(12, "student", exp=int(time.time()) - 60)
        cases = [
            ("missing JWT", "/profile", {}, 401),
            ("invalid JWT", "/profile", {"Authorization": "Bearer invalid-token"}, 401),
            ("expired JWT", "/profile", {"Authorization": "Bearer " + expired}, 401),
            ("student denied from admin endpoint", "/admin", headers(12, "student"), 403),
            ("student denied from admin notification", "/notifications/system", headers(12, "student"), 403),
        ]
        actual_rows = []
        for name, path, request_headers, expected in cases:
            if path == "/notifications/system":
                response = client.post(path, headers=request_headers, json={"title": "test", "message": "test"})
            else:
                response = client.get(path, headers=request_headers)
            actual_rows.append((name, expected, response.status_code))
            print(f"SECURITY_MATRIX case={name!r} expected_status={expected} actual_status={response.status_code}")
            assert response.status_code == expected, (
                f"{name}: expected HTTP {expected}, got HTTP {response.status_code}: {response.text}"
            )
        assert len(actual_rows) == 5
    finally:
        close_database(engine, db, tables)


def test_material_upload_idempotency_processing_listing_download_and_authorization(monkeypatch):
    engine, db, tables, resources = install_database(monkeypatch)
    try:
        client = TestClient(main.app)
        teacher = headers(11, "teacher")
        student = headers(12, "student")
        outsider = headers(13, "student")
        other_teacher = headers(14, "teacher")
        payload = b"Database indexing improves lookup performance."
        files = {"file": ("database-indexing.txt", payload, "text/plain")}

        uploaded = client.post(
            "/courses/101/resources",
            headers={**teacher, "Idempotency-Key": "material-test-key-001"},
            data={"title": "Database Indexing"},
            files=files,
        )
        assert uploaded.status_code == 200, uploaded.text
        resource_id = uploaded.json()["id"]
        assert uploaded.json()["status"] == "UPLOADED"

        # Exercise the actual extraction/process service with semantic indexing unavailable.
        assert material_service.prepare_resource(resource_id) is True
        document = resources.find_one({"_id": ObjectId(resource_id)})
        assert document["processing_status"] == "READY"
        assert document["extraction_status"] == "READY"
        assert document["indexing_status"] == "LEXICAL_ONLY"

        listed = client.get("/courses/101/resources", headers=student)
        assert listed.status_code == 200
        assert any(row["id"] == resource_id and row["processing_status"] == "READY" for row in listed.json())
        status = client.get(f"/resources/{resource_id}/status", headers=student)
        assert status.status_code == 200
        assert status.json()["status"] == "READY"
        download = client.get(f"/resources/{resource_id}/download", headers=student)
        assert download.status_code == 200
        assert download.content == payload
        assert "attachment" in download.headers.get("content-disposition", "")

        for response in (
            client.get(f"/resources/{resource_id}/status", headers=outsider),
            client.get(f"/resources/{resource_id}/download", headers=outsider),
            client.get(f"/resources/{resource_id}/status", headers=other_teacher),
            client.get(f"/resources/{resource_id}/download", headers=other_teacher),
        ):
            assert response.status_code == 403, f"expected HTTP 403; got {response.status_code}: {response.text}"

        replay = client.post(
            "/courses/101/resources",
            headers={**teacher, "Idempotency-Key": "material-test-key-001"},
            data={"title": "Database Indexing"},
            files={"file": ("database-indexing.txt", payload, "text/plain")},
        )
        assert replay.status_code == 200
        assert replay.json()["id"] == resource_id

        duplicate = client.post(
            "/courses/101/resources",
            headers=teacher,
            data={"title": "Duplicate"},
            files={"file": ("database-indexing.txt", payload, "text/plain")},
        )
        assert duplicate.status_code == 409
    finally:
        close_database(engine, db, tables)




def test_pdf_docx_and_pptx_uploads_extract_and_download_over_http(monkeypatch):
    from io import BytesIO
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    def make_pdf(text):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        ]
        output = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for number, obj in enumerate(objects, 1):
            offsets.append(len(output))
            output.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
        xref = len(output)
        output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        output.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        output.extend(
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
        )
        return bytes(output)

    docx_buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("DOCX_EXTRACT_MARKER")
    doc.save(docx_buffer)

    pptx_buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "PPTX_EXTRACT_MARKER"
    presentation.save(pptx_buffer)

    samples = [
        ("notes.pdf", "application/pdf", make_pdf("PDF_EXTRACT_MARKER"), "PDF_EXTRACT_MARKER"),
        ("notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_buffer.getvalue(), "DOCX_EXTRACT_MARKER"),
        ("notes.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", pptx_buffer.getvalue(), "PPTX_EXTRACT_MARKER"),
    ]
    engine, db, tables, resources = install_database(monkeypatch)
    try:
        client = TestClient(main.app)
        teacher = headers(11, "teacher")
        student = headers(12, "student")
        for index, (filename, content_type, payload, marker_text) in enumerate(samples, 1):
            uploaded = client.post(
                "/courses/101/resources",
                headers={**teacher, "Idempotency-Key": f"format-upload-test-{index:03d}"},
                data={"title": filename},
                files={"file": (filename, payload, content_type)},
            )
            assert uploaded.status_code == 200, f"{filename}: {uploaded.status_code} {uploaded.text}"
            resource_id = uploaded.json()["id"]
            assert material_service.prepare_resource(resource_id) is True, filename

            saved = resources.find_one({"_id": ObjectId(resource_id)})
            assert saved["processing_status"] == "READY", filename
            assert saved["extraction_status"] == "READY", filename
            assert marker_text in saved["content"], f"{filename}: extracted text was {saved.get('content')!r}"

            status = client.get(f"/resources/{resource_id}/status", headers=student)
            assert status.status_code == 200
            assert status.json()["status"] == "READY"
            listing = client.get("/courses/101/resources", headers=student)
            assert listing.status_code == 200
            assert any(row["id"] == resource_id for row in listing.json())
            downloaded = client.get(f"/resources/{resource_id}/download", headers=student)
            assert downloaded.status_code == 200
            assert downloaded.content == payload
    finally:
        close_database(engine, db, tables)


def test_upload_rejects_wrong_owner_bad_signature_and_invalid_idempotency_key(monkeypatch):
    engine, db, tables, _ = install_database(monkeypatch)
    try:
        client = TestClient(main.app)
        teacher = headers(11, "teacher")
        other_teacher = headers(14, "teacher")

        # A teacher cannot upload to a course owned by another teacher.
        forbidden_course = client.post(
            "/courses/101/resources",
            headers=other_teacher,
            data={"title": "Unauthorized"},
            files={"file": ("note.txt", b"not allowed", "text/plain")},
        )
        assert forbidden_course.status_code == 404

        # Extension alone is not sufficient: reject a PDF without a PDF signature.
        malformed = client.post(
            "/courses/101/resources",
            headers=teacher,
            data={"title": "Malformed PDF"},
            files={"file": ("malformed.pdf", b"this is not a PDF", "application/pdf")},
        )
        assert malformed.status_code == 400
        assert "Invalid PDF signature" in malformed.text

        invalid_key = client.post(
            "/courses/101/resources",
            headers={**teacher, "Idempotency-Key": "bad key with spaces"},
            data={"title": "Invalid key"},
            files={"file": ("note.txt", b"safe text", "text/plain")},
        )
        assert invalid_key.status_code == 400
        assert "Invalid Idempotency-Key" in invalid_key.text
    finally:
        close_database(engine, db, tables)



def test_assignment_submission_grading_publication_and_student_feedback_visibility(monkeypatch):
    engine, db, tables, _ = install_database(monkeypatch)
    try:
        client = TestClient(main.app)
        teacher = headers(11, "teacher")
        student = headers(12, "student")
        other_student = headers(13, "student")

        created = client.post("/assignments", headers=teacher, data={
            "course_id": "101",
            "title": "Indexing Assignment",
            "description": "Explain how a B-tree index helps database queries.",
            "type": "assignment",
            "max_marks": "100",
        })
        assert created.status_code == 200, created.text
        assignment_id = created.json()["id"]

        submitted = client.post("/submissions", headers=student, data={
            "assignment_id": str(assignment_id),
            "answer": "A B-tree index speeds up database lookups.",
        })
        assert submitted.status_code == 200, submitted.text
        submission_id = submitted.json()["id"]
        assert submitted.json()["marks"] is None
        assert submitted.json()["marks_published"] is False

        graded = client.put(
            f"/submissions/{submission_id}/marks",
            headers=teacher,
            params={"marks": 82.5, "feedback": "Good explanation of index lookup; include balancing details next time."},
        )
        assert graded.status_code == 200, graded.text
        assert graded.json()["marks_published"] is False

        before_publish = client.get("/my-marks", headers=student)
        assert before_publish.status_code == 200
        own_row = next(row for row in before_publish.json() if row["id"] == submission_id)
        assert own_row["marks"] is None
        assert own_row["teacher_review_note"] is None
        assert own_row["marks_status"] == "awaiting_publication"
        assert client.get("/my-marks", headers=other_student).json() == []

        published = client.post(f"/submissions/{submission_id}/publish", headers=teacher)
        assert published.status_code == 200, published.text
        assert published.json()["marks_published"] is True

        after_publish = client.get("/my-marks", headers=student)
        assert after_publish.status_code == 200
        own_row = next(row for row in after_publish.json() if row["id"] == submission_id)
        assert own_row["marks"] == 82.5
        assert own_row["teacher_review_note"] == "Good explanation of index lookup; include balancing details next time."
        assert own_row["marks_status"] == "published"
    finally:
        close_database(engine, db, tables)



def test_study_copilot_sources_include_supporting_evidence_and_page_references():
    from study_copilot import build_answer

    result = {
        "title": "Database Notes",
        "type": "resource",
        "course_id": 101,
        "distance": 0.12,
        "content": "[Page 3] A primary key uniquely identifies each row in a database table. "
                   "[Page 4] A foreign key links records between related tables.",
    }
    answer = build_answer("Explain primary key database", [result])
    assert answer["mode"] == "study knowledge + course context"
    source = next(source for source in answer["sources"] if source["title"] == "Database Notes")
    assert source["course_id"] == 101
    assert source["page_references"] == [3, 4]
    assert source["evidence"]
    assert any("primary key" in snippet.lower() for snippet in source["evidence"])


def test_study_copilot_source_metadata_omits_missing_page_references():
    from study_copilot import _source_payload

    source = _source_payload({
        "title": "Plain text notes",
        "type": "resource",
        "course_id": 101,
        "distance": 0.2,
        "content": "A database stores organized information.",
    }, ["A database stores organized information."])
    assert "page_references" not in source
    assert source["evidence"] == ["A database stores organized information."]
